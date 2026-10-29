# Construit le classeur Excel de l'empreinte carbone GitHub Copilot.
# Usage : python3 build.py sortie.xlsx [export.csv]
import sys, json, re, csv, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1]
CSV = sys.argv[2] if len(sys.argv) > 2 else None
SRC = json.load(open('src.json')); ECO = json.load(open('eco.json'))

# ---------- Style ----------
FONT = 'Arial'
INK, INK2, MUTED, NAVY, LINE, PANEL = '14161C', '4A4F5C', '737885', '00008F', 'E6E8ED', 'F5F6F8'
INPUT_FILL = PatternFill('solid', fgColor='FFF4CC')
PANEL_FILL = PatternFill('solid', fgColor=PANEL)
HEAD_FILL = PatternFill('solid', fgColor='EEF0F3')
WHITE = PatternFill('solid', fgColor='FFFFFF')
thin = Side(style='thin', color=LINE)
B_BOT = Border(bottom=thin)
B_BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
SER = ['2A2AD0', 'F07662', '00A3BB', 'D99A00', '914146', '7D7BE8', '3F9A5F', '9A9DB3']
OTHER = 'D2D5DC'
PROJ = 'B9BCEB'

def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

# Formats : tokens en k / M / Mds, CO2e en kg puis t
F_TOK = '[>=1000000000]#,##0.0,,," Mds";[>=1000000]#,##0.0,," M";#,##0.0," k"'
F_KG = '[>=1000]#,##0.00," t";#,##0.0" kg"'
F_KG0 = '#,##0.0'
F_KWH = '[>=1000]#,##0.0," MWh";#,##0.0" kWh"'
F_USD = '#,##0.00" $"'
F_PCT = '0.0%'
F_DPCT = '+0.0%;-0.0%;0.0%'
F_DATE = 'dd/mm/yyyy'
F_SCI = '0.00E+00'
F_INT = '#,##0'

def txt_kg(x, unit=True):
    """Texte « 12,3 kg CO₂e » ou « 1,2 t CO₂e » (séparateurs de la langue du poste)."""
    return f'IF({x}>=1000,FIXED({x}/1000,2)&" t",FIXED({x},1)&" kg")' + ('&" CO₂e"' if unit else '')

wb = Workbook()
names = {}
def name(n, ref):
    names[n] = ref
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

def setw(ws, widths):
    for k, v in widths.items(): ws.column_dimensions[k].width = v

def sheet(title, tab=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    if tab: ws.sheet_properties.tabColor = tab
    ws.column_dimensions['A'].width = 2.5
    return ws

def title(ws, t, sub=None):
    ws['B2'] = t; ws['B2'].font = f(18, True)
    ws.row_dimensions[2].height = 28
    if sub is not None:
        ws['B3'] = sub; ws['B3'].font = f(10, color=MUTED)

def section(ws, cell, t, width=1):
    ws[cell] = t; ws[cell].font = f(12, True)
    r = ws[cell].row; c = ws[cell].column
    for i in range(width):
        ws.cell(r, c + i).border = B_BOT

def head(ws, row, col, labels, wrap=True):
    for i, l in enumerate(labels):
        c = ws.cell(row, col + i, l)
        c.font = f(9, True, INK2); c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal='left' if i == 0 else 'right')
        c.border = B_BOT

def cellv(ws, ref, v, fmt=None, font=None, fill=None, align=None):
    c = ws[ref]; c.value = v
    if fmt: c.number_format = fmt
    c.font = font or f()
    if fill: c.fill = fill
    if align: c.alignment = align
    return c

def inp(ws, ref, v, fmt=None):
    return cellv(ws, ref, v, fmt, f(10, True, '0000FF'), INPUT_FILL, Alignment(horizontal='right'))

def note(ws, ref, t, size=9):
    c = ws[ref]; c.value = t; c.font = f(size, color=MUTED, italic=False)
    c.alignment = Alignment(wrap_text=False, vertical='top')
    return c

def para(ws, row, col, t, ncols, height=None, size=10, color=INK, bold=False):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + ncols - 1)
    c = ws.cell(row, col, t); c.font = f(size, bold, color)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if height: ws.row_dimensions[row].height = height
    return c

# ---------- Feuilles (ordre d'affichage) ----------
wb.remove(wb.active)
wsS = sheet('Synthèse', NAVY)
wsE = sheet('Empreinte', NAVY)
wsC = sheet('Calcul', '737885')
wsK = sheet('Contrôles', '737885')
wsM = sheet('Méthodologie', 'C9CCD4')
wsD = sheet('Données', 'D99A00')
wsI = sheet('Interne', 'E6E8ED')
wsL = sheet('Listes', 'E6E8ED')
wsH = sheet("Mode d'emploi", 'D99A00')
wb.move_sheet("Mode d'emploi", offset=-(len(wb.sheetnames) - 1))  # premier onglet

# ---------- Listes ----------
FE_SCEN = [
    ('France (2020, EEA)', 51.1, "EEA, défaut du Calculateur de Boris"),
    ('Azure, datacenters USA (2023)', 352.8, "arXiv 2505.09598, repris dans le Calculateur de Boris"),
    ('Union européenne (2025, IEA)', 170, "IEA, Electricity 2026, intensité CO₂ de la production"),
    ('Monde (2025, IEA)', 435, "IEA, Electricity 2026, intensité CO₂ de la production"),
    ('USA (2021, OWID)', 379, "Our World in Data (pas de valeur IEA publique)"),
    ('Monde (EcoLogits)', 590.478, "Mix mondial EcoLogits, Calculateur de Boris"),
    ('Personnalisé', None, "Valeur saisie dans Calcul"),
]
wsL['A1'] = 'Mix électrique'; wsL['B1'] = 'gCO₂e/kWh'; wsL['C1'] = 'Source'
for i, (n, v, s) in enumerate(FE_SCEN, start=2):
    wsL.cell(i, 1, n); wsL.cell(i, 2, v); wsL.cell(i, 3, s)
name('L_MixNom', "Listes!$A$2:$A$8"); name('L_MixVal', "Listes!$B$2:$B$8"); name('L_MixSrc', "Listes!$C$2:$C$8")
wsL['E1'] = 'Source des facteurs'; wsL['E2'] = 'Table fournie'; wsL['E3'] = 'Calculateur de Boris'
wsL['F1'] = 'Oui/Non'; wsL['F2'] = 'Oui'; wsL['F3'] = 'Non'
wsL['G1'] = 'Convention'; wsL['G2'] = 'Inclusive'; wsL['G3'] = 'Mixte'; wsL['G4'] = 'Exclusive'
wsL['I1'] = 'Mois'
MOIS = ['janv.', 'févr.', 'mars', 'avr.', 'mai', 'juin', 'juil.', 'août', 'sept.', 'oct.', 'nov.', 'déc.']
for i, m in enumerate(MOIS, start=2): wsL.cell(i, 9, m)
for c in wsL[1]: c.font = f(9, True)
setw(wsL, {'A': 30, 'B': 12, 'C': 50, 'E': 22, 'G': 14})
note(wsL, 'A11', "Listes de choix utilisées par les menus déroulants. Ne pas modifier.")

# ---------- Familles ----------
FAMS = [
    ('anthropic', 'Anthropic', 'Exclusive', "Doc Anthropic : total = cache read + cache creation + input"),
    ('openai56', 'OpenAI ≥ GPT-5.6', 'Inclusive', "Doc OpenAI : input ordinaire = input − cached − write"),
    ('openai', 'OpenAI < GPT-5.6', 'Inclusive', "Doc OpenAI : pas d'écriture facturée (ε sans objet)"),
    ('google', 'Google Gemini', 'Inclusive', "Doc Gemini : promptTokenCount inclut le cache"),
    ('xai', 'xAI Grok', 'Inclusive', "Doc xAI : le total du prompt compte le cache"),
    ('moonshot', 'Moonshot Kimi', 'Inclusive', "Doc Fireworks : cache inclus dans le total du prompt"),
    ('microsoft', 'Microsoft MAI', 'Inclusive', "Aucune source trouvée : inclusive par défaut"),
    ('other', 'Hors grille', 'Inclusive', "Modèle absent de la grille : inclusive par défaut"),
]
FAMLAB = {k: l for k, l, _, _ in FAMS}

# ---------- Modèles ----------
def norm(s):
    t = str(s or '').lower(); t = re.sub(r'^\s*auto\s*[:\-–]\s*', '', t); fast = 'fast' in t
    t = re.sub(r'\(.*?\)', ' ', t); t = re.sub(r'\bpreview\b|\bfast mode\b|\bfast\b', ' ', t)
    t = re.sub(r'(\d)[\s_-]+(?=\d)', r'\1.', t); t = re.sub(r'[^a-z0-9.]', '', t); t = re.sub(r'^claude', '', t)
    return t + ('~fast' if fast else '')

UFE = {norm(u[0]): u for u in SRC['user']}
rows = []  # dict par ligne de la table des modèles
used_fe = set()
for m in SRC['models']:
    nm, fam, p, eco = m[0], m[1], m[2], m[3]
    u = UFE.get(norm(nm)); used_fe.add(norm(nm)) if u else None
    e = ECO[eco]
    rows.append(dict(name=nm, fam=fam, p=p, before=None, eco=e, ecoexact=bool(m[4]), fe=u[1] if u else None, st=u[2] if u else None, api=len(m) > 6 and m[6]))
    if len(m) > 5 and m[5]:
        for until, ph in m[5]:
            rows.append(dict(name=nm + ' [ancien tarif]', fam=fam, p=ph, before=until, eco=e, ecoexact=bool(m[4]), fe=u[1] if u else None, st=u[2] if u else None))
for u in SRC['user']:
    if norm(u[0]) in used_fe: continue
    rows.append(dict(name=u[0], fam='other', p=None, before=None, eco=ECO['gpt54'], ecoexact=False, fe=u[1], st=u[2]))
rows.append(dict(name='(Autre modèle)', fam='other', p=None, before=None, eco=ECO['gpt54'], ecoexact=False, fe=None, st=None))
NMOD = len(rows); NBLANK = 25
M0 = 30                       # première ligne de données de la table des modèles (Calcul)
M1 = M0 + NMOD + NBLANK - 1   # dernière ligne (lignes vides pour de nouveaux modèles)

# ---------- Données ----------
EXPORT = ['date', 'username', 'product', 'sku', 'model', 'quantity', 'unit-type', 'applied_cost_per_quantity',
          'gross_amount', 'discount_amount', 'net_amount', 'total_monthly_quota', 'organization', 'app_client_id',
          'repository', 'cost_center_name', 'aic_quantity', 'aic_gross_amount', 'input', 'output', 'cache_read', 'cache_write']
def R(c): return f'TData[[#This Row],[{c}]]'
L_ = R('Ligne')
HELP = [
    ('Modele', f'IF({R("date")}="","",IF(LEFT(TRIM({R("model")}),5)="Auto:",TRIM(MID(TRIM({R("model")}),6,200)),TRIM({R("model")})))', None),
    ('Jour', f'IF({R("date")}="","",IFERROR(IF(ISNUMBER({R("date")}),INT({R("date")}),DATEVALUE({R("date")})),""))', F_DATE),
    ('Modele_cle', f'IF({R("Modele")}="","",IF(COUNTIFS(M_Nom,{R("Modele")}&" [ancien tarif]",M_Avant,">"&{R("Jour")})>0,{R("Modele")}&" [ancien tarif]",{R("Modele")}))', None),
    ('Ligne', f'IF(OR({R("Modele_cle")}="",AND({R("product")}<>"",ISERROR(SEARCH("copilot",{R("product")})))),"",IFERROR(MATCH({R("Modele_cle")},M_Nom,0),MATCH("(Autre modèle)",M_Nom,0)))', '0'),
    ('Groupe', f'IF({L_}="","",SUBSTITUTE(INDEX(M_Nom,{L_})," [ancien tarif]",""))', None),
    ('F_input_frais', f'IF({L_}="","",MAX(0,N({R("input")})-INDEX(M_Delta,{L_})*N({R("cache_read")})-INDEX(M_Eps,{L_})*N({R("cache_write")})))', F_INT),
    ('T_O', f'IF({L_}="","",{R("F_input_frais")}*INDEX(M_sI,{L_})+N({R("cache_read")})*INDEX(M_sCR,{L_})+N({R("cache_write")})*INDEX(M_sCW,{L_})+N({R("output")}))', F_INT),
    ('Energie_kWh', f'IF({L_}="","",{R("T_O")}*INDEX(M_eO,{L_})*INDEX(M_PUE,{L_})/1000)', '0.0000'),
    ('CO2e_kg', f'IF({L_}="","",{R("T_O")}*INDEX(M_FE,{L_}))', '0.0000'),
    ('Montant_brut', f'IF({R("date")}="","",IF(ISNUMBER({R("gross_amount")}),{R("gross_amount")},IFERROR(VALUE(SUBSTITUTE({R("gross_amount")},".",MID(1/2,2,1))),0)))', '0.00'),
    ('Annee', f'IF({R("Jour")}="","",YEAR({R("Jour")}))', '0'),
    ('Mois', f'IF({R("Jour")}="","",MONTH({R("Jour")}))', '0'),
    ('C2_cache_sup_input', f'IF({L_}="","",IF(N({R("cache_read")})+N({R("cache_write")})>N({R("input")}),1,0))', '0'),
    ('C4_input_frais_negatif', f'IF({L_}="","",IF(N({R("input")})-INDEX(M_Delta,{L_})*N({R("cache_read")})-INDEX(M_Eps,{L_})*N({R("cache_write")})<0,1,0))', '0'),
    ('C4_CW_sans_tarif', f'IF({L_}="","",IF(AND(N({R("cache_write")})>0,INDEX(M_PCW,{L_})="",N(INDEX(M_PO,{L_}))>0),1,0))', '0'),
]
COLS = EXPORT + [h[0] for h in HELP]
NE = len(EXPORT)
COL = {c: CL(i + 1) for i, c in enumerate(COLS)}

data = []
if CSV:
    with open(CSV, newline='', encoding='utf-8-sig') as fh:
        rd = csv.DictReader(fh)
        for r in rd:
            row = []
            for c in EXPORT:
                v = r.get(c, '')
                if c == 'date': v = dt.datetime.strptime(v, '%Y-%m-%d')
                elif v == '': v = None
                elif c in ('quantity', 'applied_cost_per_quantity', 'gross_amount', 'discount_amount', 'net_amount',
                           'total_monthly_quota', 'aic_quantity', 'aic_gross_amount'): v = float(v)
                elif c in ('input', 'output', 'cache_read', 'cache_write'): v = int(float(v))
                row.append(v)
            data.append(row)
ND = max(1, len(data))
for i, c in enumerate(COLS, start=1):
    cell = wsD.cell(1, i, c)
for ri in range(ND):
    r = ri + 2
    vals = data[ri] if data else [None] * NE
    for ci, v in enumerate(vals, start=1):
        c = wsD.cell(r, ci, v)
        if ci == 1: c.number_format = F_DATE
    for hi, (hn, fm, fmt) in enumerate(HELP):
        c = wsD.cell(r, NE + 1 + hi, '=' + fm)
        if fmt: c.number_format = fmt
tab = Table(displayName='TData', ref=f'A1:{CL(len(COLS))}{ND + 1}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
tcols = []
for i, c in enumerate(COLS, start=1):
    tc = TableColumn(id=i, name=c)
    if i > NE:
        tc.calculatedColumnFormula = TableFormula(attr_text=HELP[i - NE - 1][1])
    tcols.append(tc)
tab.tableColumns = tcols
wsD.add_table(tab)
wsD.freeze_panes = 'B2'
for i, c in enumerate(COLS, start=1):
    wsD.column_dimensions[CL(i)].width = max(11, min(22, len(c) + 3))
    h = wsD.cell(1, i); h.font = f(9, True, INK if i <= NE else NAVY)
    if i > NE:
        h.fill = PatternFill('solid', fgColor='EEF0FB')
wsD.column_dimensions['A'].width = 12
wsD.cell(1, NE + 1).comment = Comment("Colonnes calculées (bleues) : ne pas modifier. Elles se recopient seules quand le tableau s'agrandit.", 'Classeur')
for c in wsD.iter_rows(min_row=2, max_row=ND + 1):
    for x in c:
        x.font = f(9) if x.column <= NE else f(9, color=NAVY)

def dcol(n): return f"'Données'!${COL[n]}:${COL[n]}"
for nm, c in [('D_Date', 'date'), ('D_Prod', 'product'), ('D_I', 'input'), ('D_O', 'output'), ('D_CR', 'cache_read'), ('D_CW', 'cache_write'),
              ('D_Mod', 'Modele'), ('D_Jour', 'Jour'), ('D_Cle', 'Modele_cle'), ('D_Grp', 'Groupe'), ('D_TO', 'T_O'), ('D_kWh', 'Energie_kWh'),
              ('D_CO2', 'CO2e_kg'), ('D_Brut', 'Montant_brut'), ('D_An', 'Annee'), ('D_Mois', 'Mois'), ('D_C2', 'C2_cache_sup_input'),
              ('D_C4F', 'C4_input_frais_negatif'), ('D_C4CW', 'C4_CW_sans_tarif')]:
    name(nm, dcol(c))

# ---------- Calcul ----------
title(wsC, 'Calcul', 'Formules, paramètres globaux et paramètres par modèle. Les cellules jaunes sont modifiables ; tout le classeur se recalcule.')
section(wsC, 'B5', 'Réglages généraux', 4)
SET = [
    ('C6', 'Source des facteurs d\'émission', 'Table fournie', None, 'P_Source', "« Table fournie » = votre table kgCO₂e/token (23/10/2026). Les modèles absents de la table passent par le Calculateur de Boris."),
    ('C7', 'Mix électrique (Calculateur de Boris)', 'France (2020, EEA)', None, 'P_MixNom', "Sert au calcul des modèles absents de la table, et à la vue de sensibilité C6."),
    ('C8', 'Mix personnalisé (gCO₂e/kWh)', 300, '0.0', 'P_MixPerso', "Utilisé seulement si le mix choisi est « Personnalisé »."),
    ('C9', 'Mix retenu (gCO₂e/kWh)', '=IF(P_MixNom="Personnalisé",P_MixPerso,INDEX(L_MixVal,MATCH(P_MixNom,L_MixNom,0)))', '0.0', 'P_Mix', None),
    ('C10', 'Inclure la fabrication (scope 3)', 'Oui', None, 'P_Fab', "Le Calculateur de Boris ajoute la fabrication du matériel par token. Votre méthode l'exclut : mettre « Non » pour s'y conformer."),
    ('C11', "Fourchette d'incertitude basse (×)", 0.5, '0.00', 'P_Lo', "Multiplicateur appliqué au résultat central pour la borne basse."),
    ('C12', "Fourchette d'incertitude haute (×)", 2, '0.00', 'P_Hi', "Multiplicateur appliqué au résultat central pour la borne haute."),
    ('C13', 'Projection : fenêtre (jours)', 28, '0', 'P_N', "Rythme quotidien moyen des N derniers jours de données, prolongé jusqu'au 31 décembre."),
    ('C14', 'Projection : ajustement', 0, F_DPCT, 'P_Adj', "Hausse ou baisse appliquée au rythme projeté (ex. +10 % si le déploiement s'élargit)."),
]
for ref, lab, v, fmt, nm, com in SET:
    r = int(ref[1:])
    cellv(wsC, f'B{r}', lab)
    if str(v).startswith('='):
        cellv(wsC, ref, v, fmt, f(10, True), align=Alignment(horizontal='right'))
    else:
        inp(wsC, ref, v, fmt)
    if com: note(wsC, f'D{r}', com)
    name(nm, f"Calcul!${ref[0]}${r}")
dv = DataValidation(type='list', formula1='=Listes!$E$2:$E$3', allow_blank=False); wsC.add_data_validation(dv); dv.add('C6')
dv = DataValidation(type='list', formula1='=L_MixNom', allow_blank=False); wsC.add_data_validation(dv); dv.add('C7')
dv = DataValidation(type='list', formula1='=Listes!$F$2:$F$3', allow_blank=False); wsC.add_data_validation(dv); dv.add('C10')
for ref in ('C8', 'C11', 'C12', 'C13'):
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0'); wsC.add_data_validation(dv); dv.add(ref)
cellv(wsC, 'D9', '=INDEX(L_MixSrc,MATCH(P_MixNom,L_MixNom,0))', font=f(9, color=MUTED))

section(wsC, 'B16', 'Formules (par ligne de l\'export)', 4)
FX = [
    ('Input frais (tokens)', 'F = max(0, I − δ·CR − ε·CW)'),
    ('Tokens équivalent output', 'T_O = F·s_I + CR·s_CR + CW·s_CW + O'),
    ('Énergie (kWh)', 'E = T_O × e_O × PUE / 1000'),
    ('Empreinte (kgCO₂e)', 'CO₂e = T_O × FE_token'),
    ('FE_token calculé (Calculateur de Boris)', 'FE = e_O × PUE × FE_mix / 10⁶ + fab'),
]
for i, (a, b) in enumerate(FX):
    cellv(wsC, f'B{17 + i}', a); cellv(wsC, f'C{17 + i}', b, font=Font(name='Consolas', size=10, color=INK))
note(wsC, 'B23', "I, O, CR, CW : colonnes input, output, cache_read, cache_write. δ, ε : convention de l'input (1 = l'input contient déjà le cache). s_I, s_CR, s_CW : prix rapportés au prix de l'output. Les formules sont dans les colonnes bleues de l'onglet Données.")

section(wsC, f'B{M0 - 3}', 'Paramètres par modèle', 25)
note(wsC, f'B{M0 - 2}', "Grille GitHub du 20/10/2026 ($ par million de tokens). Pour un nouveau modèle, remplissez une ligne vide en bas de la table (au minimum Modèle, prix et e_O). Les noms « Auto: X » utilisent la ligne de X. « (Autre modèle) » sert aux modèles inconnus.")
MH = ['Modèle', 'Famille', 'Convention forcée', 'Convention appliquée', 'Origine', 'δ', 'ε', 'P_I', 'P_CR', 'P_CW', 'P_O', 'Tarif valable avant le',
      's_I', 's_CR', 's_CW', 'Scénario EcoLogits', 'e_O (Wh/token)', 'PUE', 'fab (kgCO₂e/token)', 'FE table (kgCO₂e/token)', 'S/H',
      'FE forcé (kgCO₂e/token)', 'FE calculé (kgCO₂e/token)', 'FE retenu (kgCO₂e/token)', 'Source du FE']
head(wsC, M0 - 1, 2, MH)
wsC.row_dimensions[M0 - 1].height = 42
# colonnes : B..Z
MC = {k: CL(2 + i) for i, k in enumerate(['Nom', 'Fam', 'Force', 'Conv', 'Orig', 'Delta', 'Eps', 'PI', 'PCR', 'PCW', 'PO', 'Avant', 'sI', 'sCR', 'sCW', 'Scen', 'eO', 'PUE', 'Fab', 'FET', 'Stat', 'FEF', 'FEC', 'FE', 'Src'])}
for k, c in MC.items():
    name('M_' + k, f"Calcul!${c}${M0}:${c}${M1}")
K0 = 62  # première ligne de la table par modèle dans Contrôles (alignée sur la table des modèles)
def kref(col, r): return f"'Contrôles'!${col}${r - M0 + K0}"
for i in range(NMOD + NBLANK):
    r = M0 + i
    d = rows[i] if i < NMOD else None
    def C(k): return f'{MC[k]}{r}'
    if d:
        cellv(wsC, C('Nom'), d['name'])
        cellv(wsC, C('Fam'), FAMLAB[d['fam']])
        if d['p']:
            for k, v in zip(['PI', 'PCR', 'PCW', 'PO'], d['p']):
                if v is not None: inp(wsC, C(k), v, '0.000')
                else: cellv(wsC, C(k), None, fill=INPUT_FILL)
        else:
            for k in ['PI', 'PCR', 'PCW', 'PO']: cellv(wsC, C(k), None, fill=INPUT_FILL)
        if d['before']: inp(wsC, C('Avant'), dt.datetime.strptime(d['before'], '%Y-%m-%d'), F_DATE)
        else: cellv(wsC, C('Avant'), None, fill=INPUT_FILL)
        cellv(wsC, C('Scen'), d['eco']['name'] + ('' if d['ecoexact'] else ' (proxy)'), fill=INPUT_FILL, font=f(10, color='0000FF'))
        inp(wsC, C('eO'), d['eco']['wh'], F_SCI)
        inp(wsC, C('PUE'), d['eco']['pue'], '0.00')
        inp(wsC, C('Fab'), d['eco']['emb'] / 1000, F_SCI)
        if d['fe'] is not None: inp(wsC, C('FET'), d['fe'], F_SCI)
        else: cellv(wsC, C('FET'), None, fill=INPUT_FILL)
        if d['st']: inp(wsC, C('Stat'), d['st'])
        else: cellv(wsC, C('Stat'), None, fill=INPUT_FILL)
    else:
        for k in ['Nom', 'Fam', 'PI', 'PCR', 'PCW', 'PO', 'Avant', 'Scen', 'eO', 'PUE', 'Fab', 'FET', 'Stat']:
            cellv(wsC, C(k), None, fill=INPUT_FILL, font=f(10, color='0000FF'))
    for k in ['Nom', 'Fam']:
        wsC[C(k)].fill = INPUT_FILL; wsC[C(k)].font = f(10, color='0000FF')
    cellv(wsC, C('Force'), None, fill=INPUT_FILL, font=f(10, color='0000FF'))
    cellv(wsC, C('FEF'), None, F_SCI, fill=INPUT_FILL, font=f(10, color='0000FF'))
    B = C('Nom')
    # convention : forçage, sinon C2 (cache > input), sinon C3 (famille), sinon défaut de la famille
    cellv(wsC, C('Conv'), f'=IF({B}="","",IF({C("Force")}<>"",{C("Force")},IF({kref("J", r)}>0,"Exclusive",IF({kref("U", r)}<>"",{kref("U", r)},IFERROR(INDEX(F_Conv,MATCH({C("Fam")},F_Fam,0)),"Inclusive")))))')
    cellv(wsC, C('Orig'), f'=IF({B}="","",IF({C("Force")}<>"","Forcée",IF({kref("J", r)}>0,"C2",IF({kref("U", r)}<>"","C3","Défaut API"))))', font=f(10, color=MUTED))
    cellv(wsC, C('Delta'), f'=IF({B}="","",IF({C("Conv")}="Exclusive",0,1))', '0')
    cellv(wsC, C('Eps'), f'=IF({B}="","",IF({C("Conv")}="Inclusive",1,0))', '0')
    PI, PCR, PCW, PO = C('PI'), C('PCR'), C('PCW'), C('PO')
    cellv(wsC, C('sI'), f'=IF({B}="","",IF(N({PO})>0,{PI}/{PO},0.2))', '0.000')
    cellv(wsC, C('sCR'), f'=IF({B}="","",IF(N({PO})>0,{PCR}/{PO},0.02))', '0.000')
    cellv(wsC, C('sCW'), f'=IF({B}="","",IF(N({PO})>0,IF({PCW}="",{PI},{PCW})/{PO},0.25))', '0.000')
    cellv(wsC, C('FEC'), f'=IF({B}="","",{C("eO")}*{C("PUE")}*P_Mix/1000000+IF(P_Fab="Oui",{C("Fab")},0))', F_SCI)
    cellv(wsC, C('FE'), f'=IF({B}="","",IF({C("FEF")}<>"",{C("FEF")},IF(AND(P_Source="Table fournie",{C("FET")}<>""),{C("FET")},{C("FEC")})))', F_SCI, f(10, True, NAVY))
    cellv(wsC, C('Src'), f'=IF({B}="","",IF({C("FEF")}<>"","Forcé",IF(AND(P_Source="Table fournie",{C("FET")}<>""),IF({C("Stat")}="H","Table (hypothèse)","Table (sourcé)"),"Calculateur de Boris")))', font=f(10, color=MUTED))
    for k in MC:
        wsC[C(k)].border = Border(bottom=Side(style='hair', color=LINE))
dv = DataValidation(type='list', formula1='=Listes!$G$2:$G$4', allow_blank=True); wsC.add_data_validation(dv); dv.add(f'{MC["Force"]}{M0}:{MC["Force"]}{M1}')
dv = DataValidation(type='list', formula1='=F_Fam', allow_blank=True); wsC.add_data_validation(dv); dv.add(f'{MC["Fam"]}{M0}:{MC["Fam"]}{M1}')
section(wsC, 'G5', 'Familles : convention par défaut (API du fournisseur)', 6)
head(wsC, 6, 7, ['Famille', 'Convention', 'Source'])
for i, (k, l, cv, src) in enumerate(FAMS):
    r = 7 + i
    cellv(wsC, f'G{r}', l); inp(wsC, f'H{r}', cv); wsC[f'H{r}'].alignment = Alignment(horizontal='left'); note(wsC, f'I{r}', src)
name('F_Fam', 'Calcul!$G$7:$G$14'); name('F_Conv', 'Calcul!$H$7:$H$14')
dv = DataValidation(type='list', formula1='=Listes!$G$2:$G$4', allow_blank=False); wsC.add_data_validation(dv); dv.add('H7:H14')
note(wsC, 'G15', "Ordre d'application : convention forcée, sinon C2 (cache > input sur une ligne : exclusive), sinon C3 (rapprochement des coûts de la famille), sinon cette valeur.")
setw(wsC, {'B': 36, 'C': 22, 'D': 14, 'E': 14, 'F': 11, 'G': 17, 'H': 12, 'I': 9, 'J': 9, 'K': 9, 'L': 9, 'M': 12, 'N': 8, 'O': 8, 'P': 8,
           'Q': 24, 'R': 12, 'S': 7, 'T': 12, 'U': 12, 'V': 6, 'W': 12, 'X': 12, 'Y': 13, 'Z': 18})
wsC.freeze_panes = None

# ---------- Interne (calculs intermédiaires) ----------
wsI['B2'] = 'Calculs intermédiaires (ne pas modifier)'; wsI['B2'].font = f(12, True)
SC = [
    ('S_Has', 'Données présentes', '=COUNT(D_Jour)>0', None),
    ('S_Last', 'Dernier jour de données', '=IF(S_Has,MAX(D_Jour),DATE(S_An,1,1))', F_DATE),
    ('S_First', 'Premier jour de données', '=IF(S_Has,MIN(D_Jour),DATE(S_An,1,1))', F_DATE),
    ('S_Cur', "Année en cours (projection active)", '=AND(S_Has,YEAR(S_Last)=S_An,S_Last<DATE(S_An,12,31))', None),
    ('S_Rate', 'Rythme projeté (kgCO₂e/jour)', '=IF(S_Has,SUMIFS(D_CO2,D_Jour,">"&(S_Last-P_N),D_Jour,"<="&S_Last)/MAX(1,P_N)*(1+P_Adj),0)', '0.000'),
    ('S_YTD', 'Émissions de l\'année (kg)', '=SUMIFS(D_CO2,D_An,S_An)', '0.000'),
    ('S_Rest', 'Jours restants projetés', '=IF(S_Cur,DATE(S_An,12,31)-S_Last,0)', '0'),
    ('S_Land', 'Atterrissage fin d\'année (kg)', '=S_YTD+S_Rate*S_Rest', '0.000'),
    ('S_Cut', 'Date de coupure comparaisons', '=IF(S_Cur,S_Last,DATE(S_An,12,31))', F_DATE),
    ('S_CurSame', 'Année, jusqu\'à la coupure (kg)', '=SUMIFS(D_CO2,D_Jour,">="&DATE(S_An,1,1),D_Jour,"<="&S_Cut)', '0.000'),
    ('S_PrevSame', 'Année précédente, même période (kg)', '=SUMIFS(D_CO2,D_Jour,">="&DATE(S_An-1,1,1),D_Jour,"<="&EDATE(S_Cut,-12))', '0.000'),
    ('S_LCM', 'Fin du dernier mois complet', '=IF(S_Cur,IF(S_Last=EOMONTH(S_Last,0),S_Last,EOMONTH(S_Last,-1)),DATE(S_An,12,31))', F_DATE),
    ('S_LCMv', 'Dernier mois complet (kg)', '=IF(S_LCM<DATE(S_An,1,1),0,SUMIFS(D_CO2,D_Jour,">"&EOMONTH(S_LCM,-1),D_Jour,"<="&S_LCM))', '0.000'),
    ('S_LCMp', 'Même mois, année précédente (kg)', '=SUMIFS(D_CO2,D_Jour,">"&EOMONTH(S_LCM,-13),D_Jour,"<="&EOMONTH(S_LCM,-12))', '0.000'),
    ('S_PrevTot', 'Année précédente complète (kg)', '=SUMIFS(D_CO2,D_An,S_An-1)', '0.000'),
    ('S_TO', 'Tokens éq. output de l\'année', '=SUMIFS(D_TO,D_An,S_An)', F_INT),
    ('S_kWh', 'Énergie de l\'année (kWh)', '=SUMIFS(D_kWh,D_An,S_An)', '0.000'),
    ('S_Brut', 'Coût brut de l\'année ($)', '=SUMIFS(D_Brut,D_An,S_An)', '0.00'),
    ('S_Lines', 'Lignes de l\'année', '=COUNTIFS(D_An,S_An)', F_INT),
]
for i, (nm, lab, fm, fmt) in enumerate(SC):
    r = 4 + i
    cellv(wsI, f'B{r}', lab); cellv(wsI, f'C{r}', fm, fmt)
    name(nm, f'Interne!$C${r}')
note(wsI, f'B{4 + len(SC) + 1}', "Agrégats par modèle (alignés ligne à ligne sur la table des modèles de Calcul) et série quotidienne de l'année affichée.")
head(wsI, M0 - 1, 2, ['Modèle', 'kgCO₂e année', 'kWh année', 'T_O année', 'Coût brut année', 'Clé de tri', 'T_O toutes années'])
for r in range(M0, M1 + 1):
    cellv(wsI, f'B{r}', f'=IF(Calcul!B{r}="","",Calcul!B{r})')
    cellv(wsI, f'C{r}', f'=IF(B{r}="",0,SUMIFS(D_CO2,D_Grp,B{r},D_An,S_An))', '0.000')
    cellv(wsI, f'D{r}', f'=IF(B{r}="",0,SUMIFS(D_kWh,D_Grp,B{r},D_An,S_An))', '0.000')
    cellv(wsI, f'E{r}', f'=IF(B{r}="",0,SUMIFS(D_TO,D_Grp,B{r},D_An,S_An))', F_INT)
    cellv(wsI, f'F{r}', f'=IF(B{r}="",0,SUMIFS(D_Brut,D_Grp,B{r},D_An,S_An))', '0.00')
    cellv(wsI, f'G{r}', f'=IF(B{r}="",-1,C{r})-ROW()*1E-12', '0.000000')
    cellv(wsI, f'H{r}', f'=IF(B{r}="",0,SUMIFS(D_TO,D_Grp,B{r}))', F_INT)
for k, c in [('I_Nom', 'B'), ('I_CO2', 'C'), ('I_kWh', 'D'), ('I_TO', 'E'), ('I_Brut', 'F'), ('I_Key', 'G')]:
    name(k, f'Interne!${c}${M0}:${c}${M1}')
# Top 8 modèles de l'année
head(wsI, 3, 10, ['Rang', 'Top modèle', 'Ligne'])
for k in range(1, 9):
    r = 3 + k
    cellv(wsI, f'J{r}', k)
    cellv(wsI, f'L{r}', f'=MATCH(LARGE(I_Key,J{r}),I_Key,0)', '0')
    cellv(wsI, f'K{r}', f'=IF(INDEX(I_CO2,L{r})>0,INDEX(I_Nom,L{r}),"")')
name('I_Top', 'Interne!$K$4:$K$11')
# Série quotidienne (365 jours)
head(wsI, M0 - 1, 10, ['Jour', 'kgCO₂e'])
D0 = M0
for k in range(365):
    r = D0 + k
    cellv(wsI, f'J{r}', f'=DATE(S_An,1,1)+{k}', F_DATE)
    cellv(wsI, f'K{r}', f'=SUMIFS(D_CO2,D_Jour,J{r})', '0.000')
setw(wsI, {'B': 36, 'C': 14, 'D': 12, 'E': 16, 'F': 14, 'G': 14, 'H': 16, 'J': 12, 'K': 26, 'L': 8})
for row in wsI.iter_rows(min_row=3, max_row=M0 + 365):
    for c in row:
        if c.value is not None and not c.font.b: c.font = f(9)

# ---------- Contrôles ----------
title(wsK, 'Contrôles', 'Contrôles de cohérence de la méthode (C1 à C6), sur toutes les données importées sauf C6 (année affichée).')
section(wsK, 'B5', 'C1 · Complétude de l\'import', 4)
C1 = [
    ('Lignes importées', '=MAX(0,COUNTA(D_Date)-1)', F_INT),
    ('Lignes avec une date lisible', '=COUNT(D_Jour)', F_INT),
    ('Lignes hors Copilot (exclues)', '=MAX(0,COUNTA(D_Prod)-1-COUNTIF(D_Prod,"*copilot*"))', F_INT),
    ('Lignes sans aucun token', '=COUNTIFS(D_I,0,D_O,0,D_CR,0,D_CW,0)', F_INT),
    ('Lignes de modèles inconnus (« Autre modèle »)', '=COUNTIF(D_Grp,"(Autre modèle)")', F_INT),
    ('Période couverte', '=IF(S_Has,TEXT(S_First,"jj/mm/aaaa")&" → "&TEXT(S_Last,"jj/mm/aaaa"),"—")', None),
]
for i, (lab, fm, fmt) in enumerate(C1):
    r = 6 + i; cellv(wsK, f'B{r}', lab); cellv(wsK, f'C{r}', fm, fmt, align=Alignment(horizontal='right'))
cellv(wsK, 'D10', '=IF(C10>0,"⚠ Ajoutez ces modèles dans Calcul (ligne vide en bas de la table) ; en attendant ils utilisent la ligne « (Autre modèle) ».","✓ Tous les modèles sont reconnus.")', font=f(9, color=MUTED))
cellv(wsK, 'D8', '=IF(C8>0,"⚠ Ces lignes ne sont pas comptées.","")', font=f(9, color=MUTED))

section(wsK, 'B14', 'C3 · Rapprochement des coûts par famille (montant brut ÷ coût recalculé)', 10)
note(wsK, 'B15', "Pour chaque hypothèse de convention, le coût est recalculé avec la grille GitHub. La bonne hypothèse donne un rapport entre 0,90 et 1,00 (remises éventuelles). Calcul agrégé par famille.")
head(wsK, 16, 2, ['Famille', 'Montant brut ($)', 'Coût inclusive ($)', 'Coût mixte ($)', 'Coût exclusive ($)', 'R inclusive', 'R mixte', 'R exclusive', 'Décision C3'])
F0 = 17
for i, (k, l, _, _) in enumerate(FAMS):
    r = F0 + i
    cellv(wsK, f'B{r}', l)
    cellv(wsK, f'C{r}', f'=SUMIF(K_Fam,B{r},K_BrutP)', F_USD)
    for col, src in zip('DEF', ['K_Ci', 'K_Cm', 'K_Ce']):
        cellv(wsK, f'{col}{r}', f'=SUMIF(K_Fam,B{r},{src})', F_USD)
    for col, cc in zip('GHI', 'DEF'):
        cellv(wsK, f'{col}{r}', f'=IF({cc}{r}>0,C{r}/{cc}{r},"")', '0.000')
    cellv(wsK, f'J{r}', f'=IF(C{r}<=0,"",IF(AND(N(G{r})>=0.895,N(G{r})<=1.005),"Inclusive",IF(AND(N(H{r})>=0.895,N(H{r})<=1.005),"Mixte",IF(AND(N(I{r})>=0.895,N(I{r})<=1.005),"Exclusive",""))))', font=f(10, True))
name('K_FamNom', f'Contrôles!$B${F0}:$B${F0 + 7}'); name('K_FamDec', f'Contrôles!$J${F0}:$J${F0 + 7}')

section(wsK, 'B27', 'C4 · Cas particuliers', 4)
cellv(wsK, 'B28', 'Lignes avec cache write non facturé (s_CW = s_I appliqué)'); cellv(wsK, 'C28', '=SUM(D_C4CW)', F_INT)
cellv(wsK, 'B29', 'Lignes où I − δ·CR − ε·CW < 0 (input frais ramené à 0)'); cellv(wsK, 'C29', '=SUM(D_C4F)', F_INT)
cellv(wsK, 'D29', '=IF(C29>0,"⚠ Vérifiez la convention des modèles concernés (colonne C4_input_frais_negatif de Données).","")', font=f(9, color=MUTED))

section(wsK, 'B31', 'C5 · Boucle globale des coûts', 4)
cellv(wsK, 'B32', 'Montant brut des modèles tarifés ($)'); cellv(wsK, 'C32', '=SUMIF(K_Cr,">0",K_BrutP)', F_USD)
cellv(wsK, 'B33', 'Coût recalculé, conventions appliquées ($)'); cellv(wsK, 'C33', '=SUM(K_Cr)', F_USD)
cellv(wsK, 'B34', 'Rapport'); cellv(wsK, 'C34', '=IF(C33>0,C32/C33,"")', '0.000', f(10, True))
cellv(wsK, 'D34', '=IF(C33<=0,"",IF(AND(C34>=0.895,C34<=1.005),"✓ Dans la plage 0,90–1,00","⚠ Hors plage 0,90–1,00 : vérifier conventions, tarifs ou remises"))', font=f(9, color=MUTED))

section(wsK, 'B36', 'C6 · Sensibilité (année affichée)', 4)
TOY = 'I_TO'
C6 = [
    ('Résultat central', '=SUM(I_CO2)'),
    ('Tout au Calculateur de Boris, mix retenu', f'=SUMPRODUCT({TOY},M_FEC)'),
    ('Calculateur de Boris, France (2020, EEA) 51,1 g/kWh', f'=SUMPRODUCT({TOY},M_eO,M_PUE)*51.1/1000000+IF(P_Fab="Oui",SUMPRODUCT({TOY},M_Fab),0)'),
    ('Calculateur de Boris, Union européenne (2025, IEA) 170 g/kWh', f'=SUMPRODUCT({TOY},M_eO,M_PUE)*170/1000000+IF(P_Fab="Oui",SUMPRODUCT({TOY},M_Fab),0)'),
    ('Calculateur de Boris, Monde (2025, IEA) 435 g/kWh', f'=SUMPRODUCT({TOY},M_eO,M_PUE)*435/1000000+IF(P_Fab="Oui",SUMPRODUCT({TOY},M_Fab),0)'),
    ('Calculateur de Boris, mix retenu, sans fabrication', f'=SUMPRODUCT({TOY},M_eO,M_PUE)*P_Mix/1000000'),
    ('Fourchette basse (résultat central × basse)', '=C37*P_Lo'),
    ('Fourchette haute (résultat central × haute)', '=C37*P_Hi'),
]
head(wsK, 36, 3, ['kgCO₂e', 'Écart / central'])
wsK['B36'].border = B_BOT
for i, (lab, fm) in enumerate(C6):
    r = 37 + i
    cellv(wsK, f'B{r}', lab); cellv(wsK, f'C{r}', fm, F_KG)
    cellv(wsK, f'D{r}', f'=IF($C$37>0,C{r}/$C$37-1,"")', F_DPCT)
wsK['C36'].alignment = Alignment(horizontal='right')

section(wsK, f'B{K0 - 3}', 'C2 et C3 par modèle (toutes données)', 20)
note(wsK, f'B{K0 - 2}', "C2 : si CR + CW dépasse I sur au moins une ligne, l'input ne contient pas le cache (convention exclusive). Coûts recalculés sans le plancher à 0 de l'input frais.")
head(wsK, K0 - 1, 2, ['Modèle', 'Famille', 'Lignes', 'Σ input', 'Σ cache read', 'Σ cache write', 'Σ output', 'Montant brut ($)',
                      'C2 : lignes cache > input', 'Coût inclusive ($)', 'Coût mixte ($)', 'Coût exclusive ($)', 'R inclusive', 'R mixte', 'R exclusive',
                      'Convention appliquée', 'Coût retenu ($)', 'R retenu', 'Verdict', 'Décision C3 (famille)'])
wsK.row_dimensions[K0 - 1].height = 42
K1 = K0 + (M1 - M0)
for i in range(M1 - M0 + 1):
    r = K0 + i; m = M0 + i
    cellv(wsK, f'B{r}', f'=IF(Calcul!B{m}="","",Calcul!B{m})')
    cellv(wsK, f'C{r}', f'=IF(B{r}="","",Calcul!C{m})')
    cellv(wsK, f'D{r}', f'=IF(B{r}="",0,COUNTIF(D_Cle,B{r}))', F_INT)
    for col, src in zip('EFGHI', ['D_I', 'D_CR', 'D_CW', 'D_O', 'D_Brut']):
        cellv(wsK, f'{col}{r}', f'=IF(B{r}="",0,SUMIF(D_Cle,B{r},{src}))', F_TOK if col != 'I' else F_USD)
    cellv(wsK, f'J{r}', f'=IF(B{r}="",0,SUMIF(D_Cle,B{r},D_C2))', F_INT)
    PI, PCR, PCW, PO = f'Calcul!{MC["PI"]}{m}', f'Calcul!{MC["PCR"]}{m}', f'Calcul!{MC["PCW"]}{m}', f'Calcul!{MC["PO"]}{m}'
    PCWe = f'IF({PCW}="",{PI},{PCW})'
    base = f'N({PCR})*F{r}+{PCWe}*G{r}+N({PO})*H{r}'
    ok = f'AND(B{r}<>"",N({PO})>0)'
    cellv(wsK, f'K{r}', f'=IF({ok},({PI}*(E{r}-F{r}-G{r})+{base})/1000000,0)', F_USD)
    cellv(wsK, f'L{r}', f'=IF({ok},({PI}*(E{r}-F{r})+{base})/1000000,0)', F_USD)
    cellv(wsK, f'M{r}', f'=IF({ok},({PI}*E{r}+{base})/1000000,0)', F_USD)
    for col, cc in zip('NOP', 'KLM'):
        cellv(wsK, f'{col}{r}', f'=IF({cc}{r}>0,I{r}/{cc}{r},"")', '0.000')
    cellv(wsK, f'Q{r}', f'=IF(B{r}="","",Calcul!{MC["Conv"]}{m})')
    cellv(wsK, f'R{r}', f'=IF({ok},({PI}*(E{r}-Calcul!{MC["Delta"]}{m}*F{r}-Calcul!{MC["Eps"]}{m}*G{r})+{base})/1000000,0)', F_USD)
    cellv(wsK, f'S{r}', f'=IF(R{r}>0,I{r}/R{r},"")', '0.000')
    cellv(wsK, f'T{r}', f'=IF(OR(B{r}="",D{r}=0),"",IF(S{r}="","Sans tarif",IF(AND(S{r}>=0.895,S{r}<=1.005),"✓ Cohérent","⚠ Hors plage")))')
    cellv(wsK, f'U{r}', f'=IF(B{r}="","",IFERROR(INDEX(K_FamDec,MATCH(C{r},K_FamNom,0)),""))')
    for c in 'BCDEFGHIJKLMNOPQRSTU': wsK[f'{c}{r}'].font = f(9)
for nm, c in [('K_Fam', 'C'), ('K_BrutP', 'I'), ('K_Ci', 'K'), ('K_Cm', 'L'), ('K_Ce', 'M'), ('K_Cr', 'R')]:
    name(nm, f"'Contrôles'!${c}${K0}:${c}${K1}")
setw(wsK, {'B': 52, 'C': 16, 'D': 14, 'E': 14, 'F': 14, 'G': 13, 'H': 13, 'I': 13, 'J': 13, 'K': 13, 'L': 13, 'M': 13, 'N': 10, 'O': 10, 'P': 10, 'Q': 13, 'R': 13, 'S': 9, 'T': 13, 'U': 14})
for row in wsK.iter_rows(min_row=5, max_row=45):
    for c in row:
        if c.value is not None and c.font.sz == 11: c.font = f(10)
# C1 « période » en texte : TEXT() dépend de la langue d'Excel → on utilise deux cellules dates
wsK['C11'] = '=S_First'; wsK['C11'].number_format = F_DATE; wsK['B11'] = 'Premier jour de données'
wsK['B12'] = 'Dernier jour de données'; wsK['C12'] = '=S_Last'; wsK['C12'].number_format = F_DATE
for c in ('B12', 'C12'): wsK[c].font = f(10)
wsK['C12'].alignment = Alignment(horizontal='right')

name('L_Mois', 'Listes!$I$2:$I$13')
def dtxt(x): return f'RIGHT("0"&DAY({x}),2)&"/"&RIGHT("0"&MONTH({x}),2)&"/"&YEAR({x})'

def style_chart(ch, w=17.5, h=7.5, fmt=None, legend=True):
    ch.width, ch.height = w, h
    ch.style = 2
    ch.x_axis.delete = False; ch.y_axis.delete = False
    from openpyxl.drawing.line import LineProperties
    from openpyxl.chart.axis import ChartLines
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill='EEF0F3')))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill='C9CCD4'))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    if fmt: ch.y_axis.number_format = fmt
    if legend: ch.legend.position = 'b'
    else: ch.legend = None
    return ch

def color_series(s, hexc, line=False, width=22000, dash=None):
    if line:
        s.graphicalProperties.line.solidFill = hexc
        s.graphicalProperties.line.width = width
        if dash: s.graphicalProperties.line.dashStyle = dash
        s.smooth = True
        s.marker.symbol = 'none'
    else:
        s.graphicalProperties.solidFill = hexc
        s.graphicalProperties.line.solidFill = 'FFFFFF'
        s.graphicalProperties.line.width = 6350

# ---------- Empreinte ----------
title(wsE, 'Empreinte', None)
cellv(wsE, 'B3', '="Année "&S_An&" · facteurs : "&P_Source&" · mix "&P_MixNom&IF(S_Cur," · projection au rythme des "&P_N&" derniers jours","")', font=f(10, color=MUTED))
section(wsE, 'B5', 'Mois par mois', 10)
head(wsE, 6, 2, ['Mois', 'Réel (kgCO₂e)', 'Projeté (kgCO₂e)', 'Total (kgCO₂e)', 'Cumul réel', 'Cumul projeté', 'Énergie', 'Tokens éq. output', 'Coût brut', 'gCO₂e par $'])
wsE.row_dimensions[6].height = 30
for m in range(1, 13):
    r = 6 + m
    cellv(wsE, f'B{r}', MOIS[m - 1])
    cellv(wsE, f'C{r}', f'=SUMIFS(D_CO2,D_An,S_An,D_Mois,{m})', F_KG0)
    cellv(wsE, f'D{r}', f'=IF(S_Cur,S_Rate*MAX(0,EOMONTH(DATE(S_An,{m},1),0)-MAX(S_Last,DATE(S_An,{m},1)-1)),0)', F_KG0, f(10, color=MUTED))
    cellv(wsE, f'E{r}', f'=C{r}+D{r}', F_KG0, f(10, True))
    cellv(wsE, f'F{r}', f'=SUM(C$7:C{r})', F_KG0)
    cellv(wsE, f'G{r}', f'=SUM(D$7:D{r})', F_KG0, f(10, color=MUTED))
    cellv(wsE, f'H{r}', f'=SUMIFS(D_kWh,D_An,S_An,D_Mois,{m})', F_KWH)
    cellv(wsE, f'I{r}', f'=SUMIFS(D_TO,D_An,S_An,D_Mois,{m})', F_TOK)
    cellv(wsE, f'J{r}', f'=SUMIFS(D_Brut,D_An,S_An,D_Mois,{m})', F_USD)
    cellv(wsE, f'K{r}', f'=IF(J{r}>0,C{r}*1000/J{r},"")', '#,##0.0')
cellv(wsE, 'B19', 'Total', font=f(10, True))
for c, fmt in zip('CDEHIJ', [F_KG0, F_KG0, F_KG0, F_KWH, F_TOK, F_USD]):
    cellv(wsE, f'{c}19', f'=SUM({c}7:{c}18)', fmt, f(10, True))
cellv(wsE, 'K19', '=IF(J19>0,C19*1000/J19,"")', '#,##0.0', f(10, True))
for c in 'BCDEFGHIJK': wsE[f'{c}19'].border = Border(top=Side(style='thin', color=INK2))
note(wsE, 'B20', "Projeté : rythme quotidien moyen des derniers jours (réglage dans Calcul), appliqué aux jours restants jusqu'au 31 décembre.")

section(wsE, 'B22', 'Par modèle (8 premiers modèles de l\'année + autres), kgCO₂e', 10)
cellv(wsE, 'B23', 'Mois', font=f(9, True, INK2), fill=HEAD_FILL)
for k in range(8):
    c = CL(3 + k)
    cellv(wsE, f'{c}23', f'=IF(INDEX(I_Top,{k + 1})="","—",INDEX(I_Top,{k + 1}))', font=f(9, True, INK2), fill=HEAD_FILL,
          align=Alignment(wrap_text=True, horizontal='right', vertical='center'))
cellv(wsE, 'K23', 'Autres', font=f(9, True, INK2), fill=HEAD_FILL, align=Alignment(horizontal='right', vertical='center'))
wsE.row_dimensions[23].height = 30
for m in range(1, 13):
    r = 23 + m
    cellv(wsE, f'B{r}', MOIS[m - 1])
    for k in range(8):
        c = CL(3 + k)
        cellv(wsE, f'{c}{r}', f'=IF(INDEX(I_Top,{k + 1})="",0,SUMIFS(D_CO2,D_Grp,INDEX(I_Top,{k + 1}),D_An,S_An,D_Mois,{m}))', F_KG0)
    cellv(wsE, f'K{r}', f'=MAX(0,C{6 + m}-SUM(C{r}:J{r}))', F_KG0)

NM = M1 - M0 + 1
R0 = 40
section(wsE, f'B{R0 - 2}', 'Classement des modèles (année affichée)', 10)
head(wsE, R0 - 1, 2, ['Rang', 'Modèle', 'kgCO₂e', 'Part', 'Énergie', 'Tokens éq. output', 'Coût brut', 'gCO₂e par $', 'FE retenu (kgCO₂e/token)', 'Source du FE'])
wsE.row_dimensions[R0 - 1].height = 30
for k in range(NM):
    r = R0 + k
    cellv(wsE, f'L{r}', f'=MATCH(LARGE(I_Key,{k + 1}),I_Key,0)', '0', f(8, color='C9CCD4'))
    cellv(wsE, f'C{r}', f'=IF(INDEX(I_CO2,L{r})>0,INDEX(I_Nom,L{r}),"")')
    cellv(wsE, f'B{r}', f'=IF(C{r}="","",{k + 1})', '0', f(10, color=MUTED))
    cellv(wsE, f'D{r}', f'=IF(C{r}="","",INDEX(I_CO2,L{r}))', F_KG0, f(10, True))
    cellv(wsE, f'E{r}', f'=IF(C{r}="","",D{r}/S_YTD)', F_PCT)
    cellv(wsE, f'F{r}', f'=IF(C{r}="","",INDEX(I_kWh,L{r}))', F_KWH)
    cellv(wsE, f'G{r}', f'=IF(C{r}="","",INDEX(I_TO,L{r}))', F_TOK)
    cellv(wsE, f'H{r}', f'=IF(C{r}="","",INDEX(I_Brut,L{r}))', F_USD)
    cellv(wsE, f'I{r}', f'=IF(C{r}="","",IF(H{r}>0,D{r}*1000/H{r},""))', '#,##0.0')
    cellv(wsE, f'J{r}', f'=IF(C{r}="","",INDEX(M_FE,L{r}))', F_SCI)
    cellv(wsE, f'K{r}', f'=IF(C{r}="","",INDEX(M_Src,L{r}))', font=f(10, color=MUTED))
wsE.conditional_formatting.add(f'E{R0}:E{R0 + NM - 1}', FormulaRule(formula=[f'AND(ISNUMBER(E{R0}),E{R0}>=0.2)'], font=Font(bold=True, color=NAVY)))
setw(wsE, {'B': 10, 'C': 30, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 15, 'J': 13, 'K': 20, 'L': 4})
wsE.column_dimensions['C'].width = 30

# Graphiques Empreinte
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 40
ch.add_data(Reference(wsE, min_col=3, max_col=11, min_row=23, max_row=35), titles_from_data=True)
ch.set_categories(Reference(wsE, min_col=2, min_row=24, max_row=35))
for s, col in zip(ch.series, SER + [OTHER]): color_series(s, col)
style_chart(ch, 20, 9, '#,##0')
ch.title = 'Émissions mensuelles par modèle (kgCO₂e)'
wsE.add_chart(ch, 'N5')

ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 40
ch.add_data(Reference(wsE, min_col=6, max_col=7, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(wsE, min_col=2, min_row=7, max_row=18))
color_series(ch.series[0], NAVY); color_series(ch.series[1], PROJ)
style_chart(ch, 20, 8, '#,##0')
ch.title = 'Cumul de l\'année, réel et projeté (kgCO₂e)'
wsE.add_chart(ch, 'N24')

ch = LineChart()
ch.add_data(Reference(wsI, min_col=11, min_row=M0 - 1, max_row=M0 + 364), titles_from_data=True)
ch.set_categories(Reference(wsI, min_col=10, min_row=M0, max_row=M0 + 364))
color_series(ch.series[0], NAVY, line=True, width=15875)
style_chart(ch, 20, 8, '#,##0.0', legend=False)
ch.x_axis.number_format = 'mmm'
ch.x_axis.tickLblSkip = 31; ch.x_axis.tickMarkSkip = 31
ch.title = 'Émissions quotidiennes (kgCO₂e)'
wsE.add_chart(ch, 'N41')

# ---------- Synthèse ----------
for c in 'BCDEFGHIJKLM': wsS.column_dimensions[c].width = 11.5
wsS['B2'] = 'Empreinte carbone de GitHub Copilot'; wsS['B2'].font = f(20, True); wsS.row_dimensions[2].height = 30
cellv(wsS, 'B3', '="Inférence des modèles utilisés via Copilot · facteurs : "&P_Source&" · méthode du 20/10/2026"', font=f(10, color=MUTED))
cellv(wsS, 'B4', f'=IF(S_Has,"Données du "&{dtxt("S_First")}&" au "&{dtxt("S_Last")}&" · "&FIXED(COUNT(D_Jour),0)&" lignes","Aucune donnée : collez votre export dans l\'onglet Données (voir Mode d\'emploi).")', font=f(10, True, INK2))
cellv(wsS, 'B6', 'Année', font=f(10, True))
inp(wsS, 'C6', '=YEAR(IF(COUNT(D_Jour)>0,MAX(D_Jour),TODAY()))', '0')
wsS['C6'].alignment = Alignment(horizontal='center')
note(wsS, 'D6', "← saisissez une autre année pour changer toute la synthèse et l'onglet Empreinte (par défaut : la dernière année des données).")
name('S_An', 'Synthèse!$C$6')
dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2100'); wsS.add_data_validation(dv); dv.add('C6')

def tile(ws, r, c0, ncol, label, value, fmt, sub, big=22):
    c1 = c0 + ncol - 1
    for rr in (r, r + 1, r + 2):
        ws.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c1)
        for cc in range(c0, c1 + 1):
            x = ws.cell(rr, cc); x.fill = PANEL_FILL
            x.border = Border(left=thin if cc == c0 else None, right=thin if cc == c1 else None,
                              top=thin if rr == r else None, bottom=thin if rr == r + 2 else None)
    a = ws.cell(r, c0, label); a.font = f(9, True, INK2); a.alignment = Alignment(indent=1, vertical='bottom')
    b = ws.cell(r + 1, c0, value); b.font = f(big, True, NAVY); b.number_format = fmt; b.alignment = Alignment(indent=1, horizontal='left', vertical='center')
    s = ws.cell(r + 2, c0, sub); s.font = f(9, color=MUTED); s.alignment = Alignment(indent=1, wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 20; ws.row_dimensions[r + 1].height = 34; ws.row_dimensions[r + 2].height = 30

tile(wsS, 8, 2, 3, '="Émissions "&S_An', '=S_YTD', F_KG, f'="Fourchette "&{txt_kg("S_YTD*P_Lo",False)}&" – "&{txt_kg("S_YTD*P_Hi",False)}')
tile(wsS, 8, 5, 3, '=IF(S_Cur,"Atterrissage fin "&S_An,"Bilan "&S_An)', '=S_Land', F_KG,
     f'=IF(S_Cur,"Déjà émis "&{txt_kg("S_YTD",False)}&" · fourchette "&{txt_kg("S_Land*P_Lo",False)}&" – "&{txt_kg("S_Land*P_Hi",False)},"Année close : émissions réelles")')
tile(wsS, 8, 8, 3, 'Tokens équivalent output', '=S_TO', F_TOK, '="Énergie "&IF(S_kWh>=1000,FIXED(S_kWh/1000,2)&" MWh",FIXED(S_kWh,1)&" kWh")&" (e_O × PUE, EcoLogits v2)"')
tile(wsS, 8, 11, 3, 'Coût brut', '=S_Brut', '#,##0" $"', '=IF(S_Brut>0,FIXED(S_YTD*1000/S_Brut,1)&" gCO₂e par dollar","Pas de montant dans l\'export")')

section(wsS, 'B12', 'Tendance', 12)
tile(wsS, 13, 2, 4, '=IF(S_Cur,S_An&" à date vs même période "&(S_An-1),S_An&" vs "&(S_An-1))', '=IF(S_PrevSame>0,S_CurSame/S_PrevSame-1,"—")', F_DPCT,
     f'={txt_kg("S_CurSame")}&" contre "&{txt_kg("S_PrevSame")}', 18)
tile(wsS, 13, 6, 4, '="Dernier mois complet ("&INDEX(L_Mois,MONTH(S_LCM))&" "&YEAR(S_LCM)&") vs "&(YEAR(S_LCM)-1)', '=IF(S_LCMp>0,S_LCMv/S_LCMp-1,"—")', F_DPCT,
     f'={txt_kg("S_LCMv")}&" contre "&{txt_kg("S_LCMp")}', 18)
tile(wsS, 13, 10, 4, '="Projection "&S_An&" vs "&(S_An-1)', '=IF(S_PrevTot>0,S_Land/S_PrevTot-1,"—")', F_DPCT,
     f'=IF(S_Cur,"Rythme des "&P_N&" derniers jours : "&FIXED(S_Rate,2)&" kg/jour","Bilan "&S_An&" contre "&{txt_kg("S_PrevTot")}&" en "&(S_An-1))', 18)
for r in (14,):
    for c in (2, 6, 10):
        wsS.cell(r, c).alignment = Alignment(indent=1, horizontal='left', vertical='center')

section(wsS, 'B17', 'Émissions mensuelles', 12)
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 35
ch.add_data(Reference(wsE, min_col=3, max_col=4, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(wsE, min_col=2, min_row=7, max_row=18))
color_series(ch.series[0], NAVY); color_series(ch.series[1], PROJ)
style_chart(ch, 23.5, 8, '#,##0')
wsS.add_chart(ch, 'B18')

TR = 36
section(wsS, f'B{TR}', 'Modèles les plus émetteurs', 12)
for c, (lab, al) in {2: ('Rang', 'center'), 3: ('Modèle', 'left'), 7: ('kgCO₂e', 'right'), 9: ('Part', 'right'), 10: ('Source du FE', 'left')}.items():
    x = wsS.cell(TR + 1, c, lab); x.font = f(9, True, INK2); x.alignment = Alignment(horizontal=al)
for c in range(2, 14): wsS.cell(TR + 1, c).fill = HEAD_FILL; wsS.cell(TR + 1, c).border = B_BOT
for k in range(5):
    r = TR + 2 + k; er = R0 + k
    cellv(wsS, f'B{r}', f'=Empreinte!B{er}', '0', f(10, color=MUTED), align=Alignment(horizontal='center'))
    wsS.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    cellv(wsS, f'C{r}', f'=Empreinte!C{er}')
    wsS.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
    cellv(wsS, f'G{r}', f'=Empreinte!D{er}', F_KG0, f(10, True, NAVY), align=Alignment(horizontal='right'))
    cellv(wsS, f'I{r}', f'=Empreinte!E{er}', F_PCT, align=Alignment(horizontal='right'))
    wsS.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
    cellv(wsS, f'J{r}', f'=Empreinte!K{er}', font=f(10, color=MUTED), align=Alignment(indent=1))
    for c in range(2, 14): wsS.cell(r, c).border = Border(bottom=Side(style='hair', color=LINE))

AR = TR + 9
section(wsS, f'B{AR}', 'À retenir', 12)
s1 = f'=IF(S_Has,"En "&S_An&", l\'usage de GitHub Copilot a émis "&{txt_kg("S_YTD")}&" (fourchette "&{txt_kg("S_YTD*P_Lo")}&" à "&{txt_kg("S_YTD*P_Hi")}&") pour "&IF(S_TO>=1000000000,FIXED(S_TO/1000000000,1)&" Mds",FIXED(S_TO/1000000,1)&" M")&" de tokens équivalent output.","Importez un export pour obtenir la synthèse.")'
s2 = '=IF(Empreinte!C40="","","Le premier modèle, "&Empreinte!C40&", représente "&FIXED(Empreinte!E40*100,0)&" % des émissions de l\'année ; les cinq premiers en représentent "&FIXED(SUM(Empreinte!E40:E44)*100,0)&" %.")'
s3 = f'=IF(NOT(S_Has),"",IF(S_Cur,"Au rythme des "&P_N&" derniers jours, l\'année se terminerait vers "&{txt_kg("S_Land")}&IF(S_PrevTot>0," ("&IF(S_Land>=S_PrevTot,"+","")&FIXED((S_Land/S_PrevTot-1)*100,0)&" % par rapport à "&(S_An-1)&").","."),"L\'année "&S_An&" est close : bilan de "&{txt_kg("S_YTD")}&"."))'
s4 = '=IF(S_Brut>0,"Intensité : "&FIXED(S_YTD*1000/S_Brut,1)&" gCO₂e par dollar facturé (montant brut).","")'
for i, s in enumerate([s1, s2, s3, s4]):
    para(wsS, AR + 1 + i, 2, s, 12, 30)

# ---------- Méthodologie ----------
title(wsM, 'Méthodologie', 'Méthode « Contrôle des tokens GitHub Copilot pour l\'empreinte carbone de l\'inférence » du 20 octobre 2026.')
wsM.column_dimensions['B'].width = 120
MT = [
    ('Objet', "Le classeur estime l'empreinte carbone de l'inférence des modèles utilisés via GitHub Copilot, jour par jour, à partir de l'export « AI usage report ». C'est la déclinaison Excel du tableau de bord HTML : mêmes formules, mêmes facteurs, mêmes contrôles."),
    ('Formules', "Input frais : F = max(0, I − δ·CR − ε·CW). Tokens équivalent output : T_O = F·s_I + CR·s_CR + CW·s_CW + O, chaque catégorie étant pondérée par son prix rapporté au prix de l'output du même modèle (s_I = P_I / P_O, etc.). Énergie : E = T_O × e_O × PUE / 1000. Empreinte : CO₂e = T_O × FE_token."),
    ('Convention de l\'input (δ, ε)', "GitHub ne documente pas si la colonne input de l'export contient déjà le cache. Pour chaque modèle : convention forcée dans Calcul, sinon C2 (si CR + CW dépasse I sur au moins une ligne, la convention est forcément exclusive), sinon C3 (rapprochement des coûts de la famille), sinon la convention documentée de l'API du fournisseur."),
    ('Prix', "Grille GitHub du 20 octobre 2026, palier standard, en dollars par million de tokens. Le palier long contexte n'est pas identifiable dans l'export. Quand le cache write n'est pas facturé mais que CW > 0, s_CW = s_I. Le prix du cache read de Claude Sonnet 5.5 était de 0,20 $ avant le 1er octobre 2026 : les lignes antérieures utilisent la ligne « [ancien tarif] ». Claude Opus 4.5, 4.6, 4.7 et Claude Sonnet 4.5 ne figurent plus dans la grille GitHub : leurs tarifs publics de l'API Anthropic sont utilisés. Les modèles « Auto: X » utilisent les paramètres de X."),
    ('Facteurs d\'émission', "FE_token provient de votre table de facteurs en kgCO₂e par token (transmise le 23 octobre 2026), dont chaque ligne est marquée S (sourcé) ou H (hypothèse). Les modèles absents de la table sont calculés avec le Calculateur de Boris (carbon-footprint-modeling-tool de Boris Ruf, scénarios EcoLogits v2, commit 50bc8b0) : FE = PUE (1,2) × énergie GPU + serveur par token × intensité du mix électrique + fabrication du matériel par token (scope 3). Les modèles de 2026 absents du Calculateur de Boris sont rattachés au scénario le plus proche (marqués « proxy »)."),
    ('Mix électrique', "Les régions d'hébergement des modèles ne sont pas publiées : le mix est traité en scénarios (France 2020 EEA 51,1 g/kWh par défaut ; UE 2025 IEA 170 ; Monde 2025 IEA 435 ; USA 2021 OWID 379 ; Azure USA 352,8 ; Monde EcoLogits 590,5). Les valeurs IEA portent sur le CO₂ seul."),
    ('Projection', "L'atterrissage de fin d'année prolonge le rythme quotidien moyen des N derniers jours de données (28 par défaut), ajusté du pourcentage choisi, jusqu'au 31 décembre. C'est une version simplifiée du modèle de projection du tableau HTML (utilisateurs × activité × usage × intensité) : Excel n'en reprend pas la fourchette P10–P90. La fourchette affichée applique au résultat central les multiplicateurs d'incertitude de Calcul (×0,5 et ×2 par défaut)."),
    ('Contrôles', "C1 complétude de l'import ; C2 cache supérieur à l'input ; C3 rapprochement des coûts par famille (rapport attendu entre 0,90 et 1,00) ; C4 cas particuliers ; C5 boucle globale des coûts ; C6 sensibilité aux facteurs et au mix. Les montants utilisés sont gross_amount (aic_gross_amount est vide dans les exports réels). Dans Excel, C3 est agrégé par famille sur toutes les lignes, alors que le tableau HTML teste les cinq lignes les plus discriminantes : la décision est la même sur l'export d'exemple."),
    ('Limites', "EcoLogits ne compte que les tokens générés : appliquer un facteur par token généré à T_O, qui intègre l'input et le cache pondérés par les prix, est une extension de méthode à déclarer. L'export ne contient pas de nombre de requêtes."),
]
r = 5
for h, t in MT:
    wsM.cell(r, 2, h).font = f(11, True); r += 1
    c = wsM.cell(r, 2, t); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical='top')
    wsM.row_dimensions[r].height = max(30, 15 * (len(t) // 135 + 1)); r += 2

# ---------- Mode d'emploi ----------
title(wsH, "Mode d'emploi", "Déclinaison Excel du tableau de bord « Empreinte carbone de GitHub Copilot ».")
wsH.column_dimensions['B'].width = 4; wsH.column_dimensions['C'].width = 110
ST = [
    ('1', "Exportez le rapport d'usage Copilot (CSV « AI usage report ») depuis GitHub."),
    ('2', "Dans l'onglet Données, videz l'ancien contenu s'il y en a : sélectionnez les lignes du tableau sous l'en-tête, clic droit › Supprimer › Lignes du tableau."),
    ('3', "Ouvrez le CSV, copiez toutes les lignes sous l'en-tête (colonnes date à cache_write, dans l'ordre de l'export) et collez-les en A2 de l'onglet Données. Le tableau s'agrandit et les colonnes bleues se calculent seules. Si les nombres décimaux sont mal lus (Excel en français), passez par Données › À partir d'un fichier texte/CSV, avec « Origine du fichier » 65001 : Unicode (UTF-8) et la virgule comme délimiteur, puis copiez le résultat."),
    ('4', "Ouvrez l'onglet Synthèse. L'année affichée est la dernière des données ; tapez une autre année dans la cellule jaune pour changer."),
    ('5', "Réglez les hypothèses dans l'onglet Calcul (source des facteurs, mix électrique, fabrication, fourchette, projection, paramètres par modèle). Vérifiez l'onglet Contrôles après chaque nouvel import, en particulier les modèles inconnus (C1)."),
]
r = 5
for n, t in ST:
    wsH.cell(r, 2, n).font = f(14, True, NAVY)
    c = wsH.cell(r, 3, t); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical='top')
    wsH.row_dimensions[r].height = 30 if len(t) < 200 else 58
    r += 1
r += 1
wsH.cell(r, 2, 'Légende').font = f(12, True); r += 1
wsH.cell(r, 2).fill = INPUT_FILL; wsH.cell(r, 3, "Cellule jaune, texte bleu : saisie ou hypothèse modifiable.").font = f(10); r += 1
wsH.cell(r, 2).fill = PatternFill('solid', fgColor='EEF0FB'); wsH.cell(r, 3, "En-tête bleu dans Données : colonne calculée, ne pas modifier.").font = f(10); r += 1
wsH.cell(r, 3, "Tout le reste est calculé par formules. Les tokens s'affichent en k, M et Mds ; les émissions en kg puis en t au-delà de 1 000 kg.").font = f(10); r += 2
wsH.cell(r, 2, "Format attendu (une ligne d'exemple, ne pas la coller)").font = f(12, True); r += 1
ex = ['2026-10-21', 'jdupont', 'copilot', 'copilot_ai_credits', 'Claude Sonnet 4.6', '168.25', 'ai-credits', '0.04', '6.73', '0', '6.73', '', 'mon-org', '', 'equipe/projet', 'Centre de coûts', '', '', '465058', '56580', '8595138', '509724']
for i, (h, v) in enumerate(zip(EXPORT, ex)):
    a = wsH.cell(r + i, 3, f'{h} : {v}' if v else f'{h} : (vide)'); a.font = Font(name='Consolas', size=9, color=INK2)
r += len(EXPORT) + 1
wsH.cell(r, 2, 'Onglets').font = f(12, True); r += 1
for t in ["Synthèse : indicateurs clés de l'année, tendance, graphique mensuel, modèles les plus émetteurs, phrases à retenir.",
          "Empreinte : tableau mois par mois (réel et projeté), répartition par modèle, cumul, série quotidienne, classement complet des modèles.",
          "Calcul : formules, réglages globaux et paramètres par modèle (prix, convention, e_O, facteurs).",
          "Contrôles : C1 à C6.", "Méthodologie : sources et choix de méthode.", "Données : l'export collé et les colonnes calculées.",
          "Interne et Listes : calculs intermédiaires et listes de choix, à ne pas modifier."]:
    wsH.cell(r, 3, '• ' + t).font = f(10); r += 1

wsI.sheet_state = 'visible'
wsL.sheet_state = 'hidden'
wb.calculation.fullCalcOnLoad = True
wb.active = wb.sheetnames.index('Synthèse')
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = (ws.title == 'Synthèse')
wb.save(OUT)
print('saved', OUT, 'rows', len(data), 'models', NMOD)
