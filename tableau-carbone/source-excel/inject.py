# Ajoute les valeurs calculées (issues du recalcul LibreOffice) au fichier openpyxl d'origine,
# pour que le classeur s'affiche rempli dès l'ouverture, sans perdre les colonnes calculées du tableau.
import sys, zipfile, re, datetime as dt
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
src, calc, out = sys.argv[1:4]
wbv = load_workbook(calc, data_only=True, read_only=True)
vals = {}
for ws in wbv.worksheets:
    d = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and hasattr(c, 'coordinate'): d[c.coordinate] = c.value
    vals[ws.title] = d
zin = zipfile.ZipFile(src)
wbxml = zin.read('xl/workbook.xml').decode()
titles = [re.search(r'name="([^"]*)"', s).group(1) for s in re.findall(r'<sheet [^>]*>', wbxml)]
titles = [t.replace('&apos;', "'").replace('&amp;', '&') for t in titles]
pat = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: s="\d+")?)((?: t="\w+")?)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)
zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
n = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    m = re.match(r'xl/worksheets/sheet(\d+)\.xml$', item.filename)
    if m:
        v = vals[titles[int(m.group(1)) - 1]]
        def rep(mm):
            global n
            ref, s, t, fm = mm.groups(); x = v.get(ref); n += 1
            if isinstance(x, bool): return f'<c r="{ref}"{s} t="b"><f>{fm}</f><v>{int(x)}</v></c>'
            if isinstance(x, (dt.datetime, dt.date)): x = to_excel(x)
            if isinstance(x, (int, float)): return f'<c r="{ref}"{s}><f>{fm}</f><v>{repr(float(x)) if isinstance(x, float) else x}</v></c>'
            if x is None: x = ''
            return f'<c r="{ref}"{s} t="str"><f>{fm}</f><v>{escape(str(x))}</v></c>'
        data = pat.sub(rep, data.decode()).encode()
    zout.writestr(item, data)
zout.close()
print('valeurs injectées :', n)
